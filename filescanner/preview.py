"""Preview pane for the Files tab: images, PDFs, video, audio, text, tables, archives and fonts.

Anything the app can't draw itself falls back to the operating system's own thumbnail
(Windows Explorer / macOS Quick Look), which covers Adobe and other formats when their apps are installed.
"""

import base64
import html
import csv
import io
import os
import re
import shutil
import subprocess
import sys
import tarfile
import tempfile
import zipfile
from xml.etree import ElementTree

from PySide6.QtCore import QObject, QRunnable, QSize, Qt, QThreadPool, QTimer, QUrl, Signal
from PySide6.QtGui import QFontDatabase, QImage, QImageReader, QMovie, QPainter, QPixmap
from PySide6.QtWidgets import (QGraphicsPixmapItem, QGraphicsScene, QGraphicsView, QHBoxLayout, QLabel,
                               QPlainTextEdit, QPushButton, QSlider, QStackedWidget, QStyle, QTableWidget,
                               QTableWidgetItem, QTextBrowser, QVBoxLayout, QWidget)

try:
    from PySide6.QtPdf import QPdfDocument
    from PySide6.QtPdfWidgets import QPdfView
    _HAS_PDF = True
except ImportError:
    _HAS_PDF = False

try:
    from PySide6.QtMultimedia import QAudioOutput, QMediaMetaData, QMediaPlayer
    from PySide6.QtMultimediaWidgets import QVideoWidget
    _HAS_MEDIA = True
except ImportError:
    _HAS_MEDIA = False

try:
    from PySide6.QtSvg import QSvgRenderer
    _HAS_SVG = True
except ImportError:
    _HAS_SVG = False

try:
    from PIL import Image, ImageOps
    _HAS_PIL = True
except ImportError:
    _HAS_PIL = False

try:
    from pillow_heif import register_heif_opener
    register_heif_opener()
except ImportError:
    pass

IS_MAC = sys.platform == "darwin"
IS_WIN = sys.platform.startswith("win")

MAX_SIDE = 1600            # largest image side kept for the zoomable view
TEXT_LIMIT = 256 * 1024    # bytes of a text file shown
TABLE_ROWS = 500
ARCHIVE_ENTRIES = 2000
SCAN_LIMIT = 96 * 1024 * 1024   # bytes read when hunting for an embedded preview

PDF_EXTS = {"pdf", "ai"}
RAW_EXTS = set("cr2 cr3 nef arw dng orf rw2 raf srw pef".split())
ADOBE_EXTS = set("ai eps psd psb indd indt".split())
TEXT_EXTS = set("txt text log md markdown html htm xhtml css js mjs cjs ts tsx jsx php py sh zsh bash command bat cmd "
                "ps1 c h cpp hpp cs java swift go rs rb kt json xml yaml yml toml ini cfg conf plist sql srt vtt ass "
                "ssa eml mbox vcf ics gltf fcpxml nfo".split())
CSV_EXTS = {"csv", "tsv"}
MARKDOWN_EXTS = {"md", "markdown"}
ARCHIVE_EXTS = set("zip jar apk tar tgz tbz2 txz".split())
COMPRESSED_TAR = (".tar.gz", ".tar.bz2", ".tar.xz")
OOXML_EXTS = set("docx docm dotx xlsx xlsm xltx pptx pptm ppsx potx".split())
ODF_EXTS = set("odt ods odp".split())
ZIP_PREVIEW_EXTS = {"pages", "numbers", "key", "xd", "sketch"}   # zip files with a preview image inside
FONT_EXTS = set("ttf otf ttc woff woff2".split())
MEDIA_CATEGORIES = {"Video", "Audio"}
NO_PREVIEW_CATEGORIES = {"Folder", "Error"}

FONT_SAMPLE = "The quick brown fox jumps over the lazy dog"


# ---------------------------------------------------------------- loaders (run in a worker thread)

def _pil_to_qimage(img):
    img = img.convert("RGBA")
    data = img.tobytes("raw", "RGBA")
    return QImage(data, img.width, img.height, img.width * 4, QImage.Format_RGBA8888).copy()


def _pil_image(source, max_side=MAX_SIDE, orientation=None):
    """Open an image with Pillow (path or bytes) and return a QImage, or None."""
    if not _HAS_PIL:
        return None
    try:
        with Image.open(io.BytesIO(source) if isinstance(source, bytes) else source) as img:
            img.draft("RGB", (max_side, max_side))
            img.load()
            if orientation:
                img.getexif()[0x0112] = orientation
            img = ImageOps.exif_transpose(img)
            img.thumbnail((max_side, max_side))
            return _pil_to_qimage(img)
    except Exception:
        return None


def _qt_image(path, max_side=MAX_SIDE):
    reader = QImageReader(path)
    reader.setAutoTransform(True)
    size = reader.size()
    if size.isValid() and max(size.width(), size.height()) > max_side:
        reader.setScaledSize(size.scaled(max_side, max_side, Qt.KeepAspectRatio))
    img = reader.read()
    return None if img.isNull() else img


def _svg_image(path, max_side=MAX_SIDE):
    if not _HAS_SVG:
        return None
    renderer = QSvgRenderer(path)
    if not renderer.isValid():
        return None
    size = renderer.defaultSize()
    if size.isEmpty():
        size = QSize(max_side, max_side)
    size = size.scaled(max_side, max_side, Qt.KeepAspectRatio)
    img = QImage(size, QImage.Format_ARGB32_Premultiplied)
    img.fill(Qt.transparent)
    painter = QPainter(img)
    renderer.render(painter)
    painter.end()
    return img


def _read_head(path, limit=SCAN_LIMIT):
    with open(path, "rb") as f:
        return f.read(limit)


def _raw_orientation(path):
    """EXIF orientation of a TIFF-based camera RAW file, if Pillow can read its header."""
    if not _HAS_PIL:
        return None
    try:
        with Image.open(path) as img:
            return img.getexif().get(0x0112)
    except Exception:
        return None


def _embedded_jpeg(data):
    """Largest JPEG embedded in a file (camera RAW previews and similar)."""
    if not _HAS_PIL:
        return None
    best, best_area = None, 0
    start = 0
    for _ in range(64):
        off = data.find(b"\xff\xd8\xff", start)
        if off < 0:
            break
        start = off + 3
        try:
            with Image.open(io.BytesIO(data[off:off + 256 * 1024])) as img:
                if img.format != "JPEG":
                    continue
                area = img.width * img.height
        except Exception:
            continue
        if area > best_area:
            best, best_area = off, area
    return data[best:] if best is not None else None


def _xmp_thumbnail(data):
    """Thumbnail stored in Adobe XMP metadata (InDesign, Illustrator, Photoshop, EPS)."""
    best = None
    for m in re.finditer(rb"xmpGImg:image(?:>|=\")([A-Za-z0-9+/=\s&#xA;]{200,}?)(?:<|\")", data):
        raw = m.group(1).replace(b"&#xA;", b"")
        raw = re.sub(rb"\s+", b"", raw)
        try:
            jpeg = base64.b64decode(raw + b"=" * (-len(raw) % 4))
        except Exception:
            continue
        if best is None or len(jpeg) > len(best):
            best = jpeg
    return _pil_image(best) if best else None


def _eps_preview(data):
    """TIFF preview inside a DOS-header EPS file."""
    if data[:4] != b"\xc5\xd0\xd3\xc6" or len(data) < 30:
        return None
    tiff_off = int.from_bytes(data[20:24], "little")
    tiff_len = int.from_bytes(data[24:28], "little")
    if tiff_off and tiff_len:
        return _pil_image(data[tiff_off:tiff_off + tiff_len])
    return None


def _ghostscript_available():
    return any(shutil.which(n) for n in ("gs", "gswin64c", "gswin32c"))


def load_image(path, ext):
    """Best effort at drawing an image-like file. Returns a QImage or None."""
    if ext == "svg":
        return _svg_image(path)
    if ext in RAW_EXTS:
        data = _read_head(path)
        jpeg = _embedded_jpeg(data)
        return _pil_image(jpeg, orientation=_raw_orientation(path)) if jpeg else None
    if ext in ADOBE_EXTS:
        img = None
        if ext in ("psd", "eps") and (ext != "eps" or _ghostscript_available()):
            img = _pil_image(path)
        if img is None:
            data = _read_head(path)
            img = (_eps_preview(data) if ext == "eps" else None) or _xmp_thumbnail(data)
        return img
    return _qt_image(path) or _pil_image(path)


def _decode_text(data):
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", "replace")
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16", "replace")
    if b"\x00" in data[:8192]:
        return None
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError as exc:
        if exc.start > len(data) - 4:    # cut mid-character at the size limit
            return data[:exc.start].decode("utf-8")
        return data.decode("cp1252" if IS_WIN else "latin-1", "replace")


def looks_like_text(path):
    try:
        with open(path, "rb") as f:
            chunk = f.read(8192)
    except OSError:
        return False
    if not chunk:
        return False
    text = _decode_text(chunk)
    if text is None:
        return False
    controls = sum(1 for ch in text if ord(ch) < 32 and ch not in "\r\n\t\f")
    return controls < len(text) * 0.02


def load_text(path, ext):
    with open(path, "rb") as f:
        data = f.read(TEXT_LIMIT + 1)
    truncated = len(data) > TEXT_LIMIT
    text = _decode_text(data[:TEXT_LIMIT])
    if text is None:
        return None
    note = f"Showing the first {TEXT_LIMIT // 1024} KB" if truncated else ""
    return ("markdown" if ext in MARKDOWN_EXTS else "text", text, note)


def load_csv(path, ext):
    with open(path, "rb") as f:
        text = _decode_text(f.read(TEXT_LIMIT)) or ""
    if ext == "tsv":
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows = []
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        rows.append(row)
        if len(rows) > TABLE_ROWS:
            break
    if not rows:
        return None
    note = f"Showing the first {TABLE_ROWS} rows" if len(rows) > TABLE_ROWS else ""
    return ("table", rows[0], rows[1:TABLE_ROWS + 1], note)


def _fmt_size(n):
    for unit in ("bytes", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:,} {unit}" if unit == "bytes" else f"{n:.1f} {unit}"
        n /= 1024


def load_archive(path):
    rows = []
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
            for i in infos[:ARCHIVE_ENTRIES]:
                d = i.date_time
                rows.append([i.filename, "" if i.is_dir() else _fmt_size(i.file_size),
                             f"{d[0]:04d}-{d[1]:02d}-{d[2]:02d} {d[3]:02d}:{d[4]:02d}"])
            total = len(infos)
    elif tarfile.is_tarfile(path):
        import time
        with tarfile.open(path) as t:
            total = 0
            for m in t:
                total += 1
                if total <= ARCHIVE_ENTRIES:
                    rows.append([m.name + ("/" if m.isdir() else ""), "" if m.isdir() else _fmt_size(m.size),
                                 time.strftime("%Y-%m-%d %H:%M", time.localtime(m.mtime))])
    else:
        return None
    note = f"{total:,} item{'s' if total != 1 else ''}" + (f", showing the first {ARCHIVE_ENTRIES:,}" if total > ARCHIVE_ENTRIES else "")
    return ("table", ["Name", "Size", "Modified"], rows, note)


def _xml_text(xml_bytes, para_tag, text_tag):
    paras = []
    for p in ElementTree.fromstring(xml_bytes).iter(para_tag):
        line = "".join(t.text or "" for t in p.iter(text_tag))
        if line.strip():
            paras.append(line)
    return paras


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_ODF_TEXT = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"


def load_office(path, ext):
    """Office / OpenDocument / iWork: a spreadsheet table, the embedded thumbnail, or a text excerpt."""
    if ext in ("xlsx", "xlsm", "xltx"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                ws = wb.worksheets[0]
                rows = []
                for row in ws.iter_rows(max_row=TABLE_ROWS + 1, max_col=50, values_only=True):
                    rows.append(["" if v is None else str(v) for v in row])
            finally:
                wb.close()
            used = max((i + 1 for row in rows for i, v in enumerate(row) if v), default=0)
            rows = [row[:used] for row in rows]
            if used:
                headers = [chr(65 + i) if i < 26 else f"{i + 1}" for i in range(len(rows[0]))]
                sheets = len(wb.sheetnames)
                note = f"Sheet “{ws.title}”" + (f" (1 of {sheets})" if sheets > 1 else "")
                if len(rows) > TABLE_ROWS:
                    note += f", first {TABLE_ROWS} rows"
                return ("table", headers, rows[:TABLE_ROWS], note)
        except Exception:
            pass
    if not zipfile.is_zipfile(path):
        return None
    with zipfile.ZipFile(path) as z:
        names = {n.lower(): n for n in z.namelist()}
        for candidate in ("docprops/thumbnail.jpeg", "docprops/thumbnail.jpg", "docprops/thumbnail.png",
                          "thumbnails/thumbnail.png", "quicklook/thumbnail.jpg", "quicklook/preview.jpg",
                          "preview.jpg", "preview-web.jpg", "preview.png", "previews/preview.png",
                          "thumbnail.png"):
            if candidate in names:
                img = _pil_image(z.read(names[candidate]))
                if img is not None:
                    return ("image", img, "Embedded thumbnail")
        paras = []
        try:
            if "word/document.xml" in names:
                paras = _xml_text(z.read(names["word/document.xml"]), _W + "p", _W + "t")
            elif ext.startswith("pp") or ext == "potx":
                slides = sorted((n for n in names if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                                key=lambda n: int(re.search(r"(\d+)\.xml$", n).group(1)))
                for i, n in enumerate(slides, 1):
                    paras.append(f"— Slide {i} —")
                    paras.extend(_xml_text(z.read(names[n]), _A + "p", _A + "t"))
            elif "content.xml" in names:
                root = ElementTree.fromstring(z.read(names["content.xml"]))
                for el in root.iter():
                    if el.tag in (_ODF_TEXT + "p", _ODF_TEXT + "h"):
                        line = "".join(el.itertext())
                        if line.strip():
                            paras.append(line)
        except ElementTree.ParseError:
            pass
    if paras:
        text = "\n\n".join(paras)
        note = "Text only"
        if len(text) > TEXT_LIMIT:
            text, note = text[:TEXT_LIMIT], f"Text only, first {TEXT_LIMIT // 1024} KB"
        return ("text", text, note)
    return None


# ---------------------------------------------------------------- operating system thumbnails

def _win_thumbnail(path, size):
    import ctypes
    from ctypes import wintypes

    class GUID(ctypes.Structure):
        _fields_ = [("Data1", ctypes.c_uint32), ("Data2", ctypes.c_uint16), ("Data3", ctypes.c_uint16),
                    ("Data4", ctypes.c_ubyte * 8)]

    class BITMAP(ctypes.Structure):
        _fields_ = [("bmType", ctypes.c_long), ("bmWidth", ctypes.c_long), ("bmHeight", ctypes.c_long),
                    ("bmWidthBytes", ctypes.c_long), ("bmPlanes", ctypes.c_ushort),
                    ("bmBitsPixel", ctypes.c_ushort), ("bmBits", ctypes.c_void_p)]

    class BITMAPINFOHEADER(ctypes.Structure):
        _fields_ = [("biSize", ctypes.c_uint32), ("biWidth", ctypes.c_long), ("biHeight", ctypes.c_long),
                    ("biPlanes", ctypes.c_ushort), ("biBitCount", ctypes.c_ushort),
                    ("biCompression", ctypes.c_uint32), ("biSizeImage", ctypes.c_uint32),
                    ("biXPelsPerMeter", ctypes.c_long), ("biYPelsPerMeter", ctypes.c_long),
                    ("biClrUsed", ctypes.c_uint32), ("biClrImportant", ctypes.c_uint32)]

    ole32, shell32 = ctypes.windll.ole32, ctypes.windll.shell32
    gdi32, user32 = ctypes.windll.gdi32, ctypes.windll.user32
    ole32.CoInitializeEx.argtypes = [ctypes.c_void_p, ctypes.c_uint32]
    shell32.SHCreateItemFromParsingName.argtypes = [ctypes.c_wchar_p, ctypes.c_void_p, ctypes.POINTER(GUID),
                                                    ctypes.POINTER(ctypes.c_void_p)]
    shell32.SHCreateItemFromParsingName.restype = ctypes.c_long
    gdi32.GetObjectW.argtypes = [ctypes.c_void_p, ctypes.c_int, ctypes.c_void_p]
    gdi32.GetDIBits.argtypes = [ctypes.c_void_p, ctypes.c_void_p, ctypes.c_uint, ctypes.c_uint, ctypes.c_void_p,
                                ctypes.c_void_p, ctypes.c_uint]
    gdi32.DeleteObject.argtypes = [ctypes.c_void_p]
    user32.GetDC.argtypes = [ctypes.c_void_p]
    user32.GetDC.restype = ctypes.c_void_p
    user32.ReleaseDC.argtypes = [ctypes.c_void_p, ctypes.c_void_p]

    # IShellItemImageFactory {BCC18B79-BA16-442F-80C4-8A59C30C463B}
    iid = GUID(0xBCC18B79, 0xBA16, 0x442F, (ctypes.c_ubyte * 8)(0x80, 0xC4, 0x8A, 0x59, 0xC3, 0x0C, 0x46, 0x3B))
    initialised = ole32.CoInitializeEx(None, 0x2) in (0, 1)   # apartment threaded; S_OK or S_FALSE
    factory = ctypes.c_void_p()
    try:
        if shell32.SHCreateItemFromParsingName(os.path.normpath(os.path.abspath(path)), None, ctypes.byref(iid),
                                               ctypes.byref(factory)) != 0 or not factory:
            return None
        vtable = ctypes.cast(ctypes.cast(factory, ctypes.POINTER(ctypes.c_void_p))[0],
                             ctypes.POINTER(ctypes.c_void_p))
        release = ctypes.WINFUNCTYPE(ctypes.c_ulong, ctypes.c_void_p)(vtable[2])
        get_image = ctypes.WINFUNCTYPE(ctypes.c_long, ctypes.c_void_p, wintypes.SIZE, ctypes.c_int,
                                       ctypes.POINTER(ctypes.c_void_p))(vtable[3])
        hbitmap = ctypes.c_void_p()
        # SIIGBF_BIGGERSIZEOK | SIIGBF_THUMBNAILONLY: a real thumbnail, never just the file-type icon
        hr = get_image(factory, wintypes.SIZE(size, size), 0x1 | 0x8, ctypes.byref(hbitmap))
        release(factory)
        if hr != 0 or not hbitmap:
            return None
        try:
            bm = BITMAP()
            gdi32.GetObjectW(hbitmap, ctypes.sizeof(bm), ctypes.byref(bm))
            w, h = bm.bmWidth, abs(bm.bmHeight)
            if w <= 0 or h <= 0:
                return None
            header = BITMAPINFOHEADER(ctypes.sizeof(BITMAPINFOHEADER), w, -h, 1, 32, 0, 0, 0, 0, 0, 0)
            buf = ctypes.create_string_buffer(w * h * 4)
            hdc = user32.GetDC(None)
            lines = gdi32.GetDIBits(hdc, hbitmap, 0, h, buf, ctypes.byref(header), 0)
            user32.ReleaseDC(None, hdc)
            if lines != h:
                return None
            data = buf.raw
            fmt = QImage.Format_ARGB32_Premultiplied if any(data[3::4]) else QImage.Format_RGB32
            return QImage(data, w, h, w * 4, fmt).copy()
        finally:
            gdi32.DeleteObject(hbitmap)
    finally:
        if initialised:
            ole32.CoUninitialize()


def _mac_thumbnail(path, size):
    out = tempfile.mkdtemp(prefix="filescanner-ql-")
    try:
        subprocess.run(["/usr/bin/qlmanage", "-t", "-s", str(size), "-o", out, path],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=15)
        for name in os.listdir(out):
            img = QImage(os.path.join(out, name))
            if not img.isNull():
                return img
        return None
    except (OSError, subprocess.SubprocessError):
        return None
    finally:
        shutil.rmtree(out, ignore_errors=True)


def system_thumbnail(path, size=1024):
    """The thumbnail Windows Explorer or macOS Quick Look would show, or None."""
    try:
        if IS_WIN:
            return _win_thumbnail(path, size)
        if IS_MAC:
            return _mac_thumbnail(path, size)
    except Exception:
        return None
    return None


# ---------------------------------------------------------------- deciding what to show

def build_preview(path, ext, category, system_first=False):
    """Work out the preview for a file. Runs in a worker thread.

    Returns one of:
      ("image", QImage, note)
      ("text" | "markdown", str, note)
      ("table", headers, rows, note)
      ("message", str)
    """
    lower = path.lower()
    if ext == "pdf":   # the PDF viewer couldn't open it
        img = system_thumbnail(path)
        return ("image", img, "System preview") if img is not None else ("message", "Couldn't read this PDF.")
    if ext in CSV_EXTS:
        return load_csv(path, ext) or ("message", "This file is empty.")
    if ext in ARCHIVE_EXTS or lower.endswith(COMPRESSED_TAR):
        result = load_archive(path)
        if result:
            return result
    if ext in OOXML_EXTS or ext in ODF_EXTS or ext in ZIP_PREVIEW_EXTS:
        result = load_office(path, ext)
        if result:
            return result
    if ext in TEXT_EXTS and ext != "svg":
        result = load_text(path, ext)
        if result:
            return result

    if system_first:
        img = system_thumbnail(path)
        if img is not None:
            return ("image", img, "System preview")

    if category == "Image" or ext in RAW_EXTS or ext in ADOBE_EXTS:
        img = load_image(path, ext)
        if img is not None:
            note = "Embedded preview" if ext in RAW_EXTS or ext in ("indd", "indt") else ""
            return ("image", img, note)
    elif category not in MEDIA_CATEGORIES and looks_like_text(path):
        result = load_text(path, ext)
        if result:
            return result

    if not system_first:
        img = system_thumbnail(path)
        if img is not None:
            return ("image", img, "System preview")
    if ext == "svg":
        result = load_text(path, ext)
        if result:
            return result
    if category == "Video":
        return ("message", "This video format can't be played here, and the system has no preview for it.")
    if ext == "eps" and not _ghostscript_available():
        return ("message", "No preview: this EPS has no built-in preview image. Installing Ghostscript adds "
                           "support for drawing EPS files.")
    return ("message", "No preview available for this file type.")


class _Signals(QObject):
    done = Signal(int, object)


class _PreviewJob(QRunnable):
    def __init__(self, signals, request, path, ext, category, system_first):
        super().__init__()
        self.signals, self.request = signals, request
        self.args = (path, ext, category, system_first)

    def run(self):
        try:
            result = build_preview(*self.args)
        except Exception as exc:
            result = ("message", f"Couldn't preview this file: {exc}")
        self.signals.done.emit(self.request, result)


# ---------------------------------------------------------------- views

class ImageView(QGraphicsView):
    """Zoomable image. Scroll to zoom, drag to pan, double-click to fit."""

    def __init__(self):
        super().__init__()
        self.setScene(QGraphicsScene(self))
        self.item = QGraphicsPixmapItem()
        self.item.setTransformationMode(Qt.SmoothTransformation)
        self.scene().addItem(self.item)
        self.setDragMode(QGraphicsView.ScrollHandDrag)
        self.setTransformationAnchor(QGraphicsView.AnchorUnderMouse)
        self.setRenderHints(QPainter.Antialiasing | QPainter.SmoothPixmapTransform)
        self.setFrameShape(QGraphicsView.NoFrame)
        self.fit = True
        self.movie = None

    def set_image(self, image):
        self._stop_movie()
        self._set_pixmap(QPixmap.fromImage(image))

    def set_movie(self, path):
        self._stop_movie()
        self.movie = QMovie(path)
        self.movie.frameChanged.connect(lambda _: self.movie and self._set_pixmap(self.movie.currentPixmap(), refit=False))
        self.movie.start()
        if self.movie.isValid():
            self._set_pixmap(self.movie.currentPixmap())
            return True
        self._stop_movie()
        return False

    def _set_pixmap(self, pix, refit=True):
        self.item.setPixmap(pix)
        self.scene().setSceneRect(self.item.boundingRect())
        if refit:
            self.fit = True
        if self.fit:
            self._fit()

    def _stop_movie(self):
        if self.movie:
            self.movie.stop()
            self.movie.deleteLater()
            self.movie = None

    def clear(self):
        self._stop_movie()
        self.item.setPixmap(QPixmap())

    def _fit(self):
        if self.item.pixmap().isNull():
            return
        self.resetTransform()
        rect = self.item.boundingRect()
        view = self.viewport().rect()
        # never enlarge small images past 100% when fitting
        if rect.width() > view.width() or rect.height() > view.height():
            self.fitInView(self.item, Qt.KeepAspectRatio)

    def resizeEvent(self, e):
        super().resizeEvent(e)
        if self.fit:
            self._fit()

    def wheelEvent(self, e):
        steps = e.angleDelta().y() / 120
        if not steps:
            return
        self.fit = False
        self.scale(1.2 ** steps, 1.2 ** steps)

    def mouseDoubleClickEvent(self, e):
        if self.fit:
            self.fit = False
            self.resetTransform()
        else:
            self.fit = True
            self._fit()


class PdfView(QWidget):
    def __init__(self):
        super().__init__()
        self.doc = QPdfDocument(self)
        self.view = QPdfView()
        self.view.setDocument(self.doc)
        self.view.setPageMode(QPdfView.PageMode.MultiPage)
        self.view.setZoomMode(QPdfView.ZoomMode.FitToWidth)
        self.pages = QLabel()
        self.fit_btn = QPushButton("Fit page")
        self.fit_btn.setCheckable(True)
        self.fit_btn.toggled.connect(lambda on: self.view.setZoomMode(
            QPdfView.ZoomMode.FitInView if on else QPdfView.ZoomMode.FitToWidth))
        self.view.pageNavigator().currentPageChanged.connect(self._update_pages)
        bar = QHBoxLayout()
        bar.setContentsMargins(4, 2, 4, 2)
        bar.addWidget(self.pages, 1)
        bar.addWidget(self.fit_btn)
        v = QVBoxLayout(self)
        v.setContentsMargins(0, 0, 0, 0)
        v.setSpacing(0)
        v.addWidget(self.view, 1)
        v.addLayout(bar)

    def load(self, path):
        """Returns None on success, or a message explaining why the PDF can't be shown."""
        self.doc.close()
        err = self.doc.load(path)
        if err != QPdfDocument.Error.None_:
            self.doc.close()
            if err == QPdfDocument.Error.IncorrectPassword:
                return "This PDF is password protected."
            return "Couldn't read this PDF."
        self.view.pageNavigator().jump(0, self.view.pageNavigator().currentLocation(), 0)
        self._update_pages()
        return None

    def _update_pages(self, *_):
        count = self.doc.pageCount()
        page = self.view.pageNavigator().currentPage() + 1
        self.pages.setText(f"Page {page} of {count}" if count else "")

    def clear(self):
        self.doc.close()
        self.pages.clear()


def _fmt_ms(ms):
    s = max(0, ms) // 1000
    return f"{s // 3600}:{s // 60 % 60:02d}:{s % 60:02d}" if s >= 3600 else f"{s // 60}:{s % 60:02d}"


class MediaView(QWidget):
    failed = Signal(str)

    def __init__(self):
        super().__init__()
        self.player = QMediaPlayer(self)
        self.audio = QAudioOutput(self)
        self.audio.setVolume(0.8)
        self.player.setAudioOutput(self.audio)
        self.video = QVideoWidget()
        self.player.setVideoOutput(self.video)
        self.cover = QLabel()
        self.cover.setAlignment(Qt.AlignCenter)
        self.cover.setMinimumHeight(120)

        style = self.style()
        self.play_btn = QPushButton()
        self.play_btn.setIcon(style.standardIcon(QStyle.SP_MediaPlay))
        self.play_btn.clicked.connect(self._toggle)
        self.slider = QSlider(Qt.Horizontal)
        self.slider.sliderMoved.connect(self.player.setPosition)
        self.time = QLabel("0:00 / 0:00")
        self.mute_btn = QPushButton()
        self.mute_btn.setCheckable(True)
        self.mute_btn.setIcon(style.standardIcon(QStyle.SP_MediaVolume))
        self.mute_btn.toggled.connect(self._mute)
        self.volume = QSlider(Qt.Horizontal)
        self.volume.setRange(0, 100)
        self.volume.setValue(80)
        self.volume.setMaximumWidth(80)
        self.volume.valueChanged.connect(lambda v: self.audio.setVolume(v / 100))

        bar = QHBoxLayout()
        bar.setContentsMargins(4, 2, 4, 2)
        bar.addWidget(self.play_btn)
        bar.addWidget(self.slider, 1)
        bar.addWidget(self.time)
        bar.addWidget(self.mute_btn)
        bar.addWidget(self.volume)
        v = QVBoxLayout(self)
        v.setContentsMargins(0, 0, 0, 0)
        v.setSpacing(0)
        v.addWidget(self.video, 1)
        v.addWidget(self.cover, 1)
        v.addLayout(bar)

        self.priming = False
        self.is_video = False
        self.player.durationChanged.connect(self._duration)
        self.player.positionChanged.connect(self._position)
        self.player.playbackStateChanged.connect(self._state)
        self.player.mediaStatusChanged.connect(self._status)
        self.player.errorOccurred.connect(self._error)
        self.player.metaDataChanged.connect(self._metadata)
        self.video.videoSink().videoFrameChanged.connect(self._frame)

    def load(self, path, is_video):
        self.clear()
        self.is_video = is_video
        self.video.setVisible(is_video)
        self.cover.setVisible(not is_video)
        self.cover.setText("♫")
        self.cover.setStyleSheet("font-size: 64px; color: palette(mid);")
        self.player.setSource(QUrl.fromLocalFile(path))

    def clear(self):
        self.priming = False
        self.player.stop()
        self.player.setSource(QUrl())   # releases the file so it can be moved
        self.slider.setRange(0, 0)
        self.time.setText("0:00 / 0:00")

    def _toggle(self):
        if self.player.playbackState() == QMediaPlayer.PlayingState:
            self.player.pause()
        else:
            self._end_priming()
            if self.player.mediaStatus() == QMediaPlayer.EndOfMedia:
                self.player.setPosition(0)
            self.player.play()

    def _mute(self, on):
        self.audio.setMuted(on or self.priming)
        self.mute_btn.setIcon(self.style().standardIcon(QStyle.SP_MediaVolumeMuted if on else QStyle.SP_MediaVolume))

    def _status(self, status):
        # show the first frame of a video without playing it: start muted, pause on the first frame
        if status == QMediaPlayer.LoadedMedia and self.is_video and not self.priming \
                and self.player.playbackState() == QMediaPlayer.StoppedState:
            self.priming = True
            self.audio.setMuted(True)
            self.player.play()
            QTimer.singleShot(3000, self._priming_timeout)

    def _priming_timeout(self):
        if self.priming:    # no frame arrived: don't leave it playing silently
            self.player.pause()
            self.player.setPosition(0)
            self._end_priming()

    def _frame(self, frame):
        if self.priming and frame.isValid():
            self.player.pause()
            self.player.setPosition(0)
            self._end_priming()

    def _end_priming(self):
        if self.priming:
            self.priming = False
            self.audio.setMuted(self.mute_btn.isChecked())

    def _error(self, _error, message):
        self.clear()
        self.failed.emit(message or "This format can't be played.")

    def _metadata(self):
        if self.is_video:
            return
        meta = self.player.metaData()
        for key in (QMediaMetaData.CoverArtImage, QMediaMetaData.ThumbnailImage):
            img = meta.value(key)
            if isinstance(img, QImage) and not img.isNull():
                self.cover.setStyleSheet("")
                self.cover.setPixmap(QPixmap.fromImage(img).scaled(
                    QSize(320, 320), Qt.KeepAspectRatio, Qt.SmoothTransformation))
                return

    def _duration(self, ms):
        self.slider.setRange(0, ms)
        self._position(self.player.position())

    def _position(self, ms):
        if not self.slider.isSliderDown():
            self.slider.setValue(ms)
        self.time.setText(f"{_fmt_ms(ms)} / {_fmt_ms(self.player.duration())}")

    def _state(self, state):
        icon = QStyle.SP_MediaPause if state == QMediaPlayer.PlayingState and not self.priming else QStyle.SP_MediaPlay
        self.play_btn.setIcon(self.style().standardIcon(icon))


class FontView(QTextBrowser):
    def __init__(self):
        super().__init__()
        self.font_id = -1

    def load(self, path):
        self.clear()
        self.font_id = QFontDatabase.addApplicationFont(path)
        families = QFontDatabase.applicationFontFamilies(self.font_id) if self.font_id >= 0 else []
        if not families:
            self.clear()
            return False
        family = families[0]
        css_family = html.escape(family.replace("'", "\\'"))
        parts = [f"<p><b>{html.escape(family)}</b>" + (f" (+{len(families) - 1} more in this file)" if len(families) > 1 else "")
                 + "</p>"]
        parts.append(f"<p style=\"font-family:'{css_family}'; font-size:20pt\">"
                     "ABCDEFGHIJKLMNOPQRSTUVWXYZ<br>abcdefghijklmnopqrstuvwxyz<br>0123456789 !?&amp;@£$%()</p>")
        for size in (12, 18, 24, 36, 48):
            parts.append(f"<p style=\"font-family:'{css_family}'; font-size:{size}pt\">{FONT_SAMPLE}</p>")
        self.setHtml("".join(parts))
        return True

    def clear(self):
        super().clear()
        if self.font_id >= 0:
            QFontDatabase.removeApplicationFont(self.font_id)
            self.font_id = -1


class TableView(QTableWidget):
    def load(self, headers, rows):
        self.clear()
        width = max([len(headers)] + [len(r) for r in rows]) if rows or headers else 0
        self.setColumnCount(width)
        self.setRowCount(len(rows))
        self.setHorizontalHeaderLabels(list(headers) + [""] * (width - len(headers)))
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                self.setItem(r, c, QTableWidgetItem(value))
        self.resizeColumnsToContents()
        for c in range(width):
            self.setColumnWidth(c, min(self.columnWidth(c), 300))


# ---------------------------------------------------------------- the pane

class PreviewPane(QWidget):
    """Shows a preview of the selected file. Call show_record() with a scan record, or clear()."""

    def __init__(self):
        super().__init__()
        self.stack = QStackedWidget()
        self.message = QLabel()
        self.message.setAlignment(Qt.AlignCenter)
        self.message.setWordWrap(True)
        self.message.setStyleSheet("color: palette(mid); padding: 12px;")
        self.image = ImageView()
        self.text = QPlainTextEdit()
        self.text.setReadOnly(True)
        self.text.setLineWrapMode(QPlainTextEdit.NoWrap)
        mono = QFontDatabase.systemFont(QFontDatabase.FixedFont)
        self.text.setFont(mono)
        self.markdown = QTextBrowser()
        self.markdown.setOpenExternalLinks(True)
        self.table = TableView()
        self.table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.table.verticalHeader().setDefaultSectionSize(20)
        self.font_view = FontView()
        self.pdf = PdfView() if _HAS_PDF else None
        self.media = MediaView() if _HAS_MEDIA else None
        if self.media:
            self.media.failed.connect(self._media_failed)
        for w in (self.message, self.image, self.text, self.markdown, self.table, self.font_view, self.pdf,
                  self.media):
            if w is not None:
                self.stack.addWidget(w)

        self.note = QLabel()
        self.note.setStyleSheet("color: palette(mid); font-size: 11px; padding: 2px 4px;")
        v = QVBoxLayout(self)
        v.setContentsMargins(0, 0, 0, 0)
        v.setSpacing(0)
        v.addWidget(self.stack, 1)
        v.addWidget(self.note)

        self.pool = QThreadPool(self)
        self.pool.setMaxThreadCount(2)
        self.signals = _Signals(self)
        self.signals.done.connect(self._job_done)
        self.request = 0
        self.pending = None
        self.record = None
        self.setMinimumHeight(160)
        self.hide()

    # ---- public

    def show_record(self, rec):
        self._reset()
        self.record = rec
        path, category = rec.get("path", ""), rec.get("category", "")
        ext = (rec.get("extension") or "").lower()
        if not path or category in NO_PREVIEW_CATEGORIES or not os.path.isfile(path):
            self.hide()
            return

        if ext in PDF_EXTS and self.pdf and self._is_pdf(path):
            error = self.pdf.load(path)
            if error is None:
                return self._show(self.pdf, "Illustrator file (PDF compatible)" if ext == "ai" else "")
            if ext == "pdf":
                return self._show_message(error)
        if category in MEDIA_CATEGORIES and self.media and ext not in ("mid", "midi"):
            self.media.load(path, category == "Video")
            return self._show(self.media)
        if ext in FONT_EXTS and self.font_view.load(path):
            return self._show(self.font_view)
        if ext == "gif" and self.image.set_movie(path):
            return self._show(self.image)
        self._start_job(path, ext, category)

    def clear(self):
        """Stop playback, release any open file (call before files are moved) and hide the pane."""
        self._reset()
        self.hide()

    def _reset(self):
        self.request += 1
        self.pool.clear()
        self.record = None
        self.image.clear()
        self.text.clear()
        self.markdown.clear()
        self.table.clear()
        self.font_view.clear()
        if self.pdf:
            self.pdf.clear()
        if self.media:
            self.media.clear()
        self.note.clear()

    def shutdown(self):
        self.clear()
        self.pool.waitForDone(3000)

    # ---- internals

    @staticmethod
    def _is_pdf(path):
        try:
            with open(path, "rb") as f:
                return b"%PDF" in f.read(1024)
        except OSError:
            return False

    def _start_job(self, path, ext, category, system_first=None):
        if system_first is None:
            # camera RAW: the OS knows the orientation and colour; 3D/CAD and design files: only the OS can draw them
            system_first = ext in RAW_EXTS or ext == "rtf" or category in ("3D / CAD", "Design")
        req = self.pending = self.request
        QTimer.singleShot(250, lambda: req == self.pending and self._show_message("Loading preview…"))
        self._show_message("")
        self.pool.start(_PreviewJob(self.signals, req, path, ext, category, system_first))

    def _job_done(self, request, result):
        if request != self.request:
            return
        self.pending = None
        kind = result[0]
        if kind == "image":
            self.image.set_image(result[1])
            source = f"{result[2]} · " if result[2] else ""
            self._show(self.image, f"{source}Scroll to zoom, double-click to fit or show actual size")
        elif kind == "text":
            self.text.setPlainText(result[1])
            self._show(self.text, result[2])
        elif kind == "markdown":
            self.markdown.setMarkdown(result[1])
            self._show(self.markdown, result[2])
        elif kind == "table":
            self.table.load(result[1], result[2])
            self._show(self.table, result[3])
        else:
            self._show_message(result[1])

    def _media_failed(self, message):
        rec = self.record
        if not rec:
            return
        if rec.get("category") == "Video":
            # the player can't decode it: fall back to the OS thumbnail of the video
            self._start_job(rec["path"], (rec.get("extension") or "").lower(), "Video", system_first=True)
        else:
            self._show_message(f"Can't play this file: {message}")

    def _show(self, widget, note=""):
        self.stack.setCurrentWidget(widget)
        self.note.setText(note)
        self.note.setVisible(bool(note))
        self.show()

    def _show_message(self, text):
        self.message.setText(text)
        self._show(self.message)
