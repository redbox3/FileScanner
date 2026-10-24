"""Export scan results to CSV, Excel, JSON or an HTML report."""

import csv
import html
import json
import os
from collections import Counter, defaultdict
from datetime import datetime

from .scanner import human_size

# (key, column heading, ticked by default)
COLUMNS = [
    ("name", "File name", True),
    ("extension", "Extension", True),
    ("category", "Category", True),
    ("type", "File type", True),
    ("mime", "MIME type", False),
    ("apps", "Probable apps (this computer)", True),
    ("apps_mac", "Probable apps (Mac)", False),
    ("apps_windows", "Probable apps (Windows)", False),
    ("default_app", "Default app (Windows)", False),
    ("size", "Size (bytes)", True),
    ("size_human", "Size", True),
    ("path", "Full path", True),
    ("folder", "Folder", False),
    ("relative_path", "Relative path", False),
    ("depth", "Folder depth", False),
    ("created", "Created", True),
    ("modified", "Modified", True),
    ("accessed", "Last opened", False),
    ("hidden", "Hidden", True),
    ("system", "System file", False),
    ("read_only", "Read only", False),
    ("symlink", "Symlink", False),
    ("width", "Width (px)", True),
    ("height", "Height (px)", True),
    ("resolution", "Resolution", False),
    ("megapixels", "Megapixels", False),
    ("orientation", "Orientation", False),
    ("duration", "Length (seconds)", False),
    ("duration_human", "Length", True),
    ("frame_rate", "Frame rate", False),
    ("hash", "Checksum (BLAKE2b)", False),
    ("duplicate_group", "Duplicate group", True),
    ("metadata", "Metadata / EXIF", True),
    ("error", "Errors", False),
]
LABELS = {k: label for k, label, _ in COLUMNS}

FORMATS = {
    "CSV (Excel, Numbers, Google Sheets)": ".csv",
    "Excel workbook (.xlsx, with summary + duplicates sheets)": ".xlsx",
    "JSON (for other software)": ".json",
    "HTML report (open in a browser)": ".html",
    "Tab separated text": ".tsv",
}


def _metadata_text(meta):
    return "; ".join(f"{k}: {v}" for k, v in (meta or {}).items())


def _rows(records, columns, split_metadata):
    """Yield (headers, list of row lists)."""
    meta_keys = []
    if split_metadata and "metadata" in columns:
        seen = {}
        for r in records:
            for k in (r.get("metadata") or {}):
                seen.setdefault(k, None)
        meta_keys = list(seen)
    headers = []
    for c in columns:
        if c == "metadata" and split_metadata:
            headers.extend(meta_keys)
        else:
            headers.append(LABELS.get(c, c))
    rows = []
    for r in records:
        row = []
        for c in columns:
            if c == "metadata":
                meta = r.get("metadata") or {}
                if split_metadata:
                    row.extend(meta.get(k, "") for k in meta_keys)
                else:
                    row.append(_metadata_text(meta))
            else:
                v = r.get(c, "")
                if isinstance(v, bool):
                    v = "Yes" if v else "No"
                row.append("" if v is None else v)
        rows.append(row)
    return headers, rows


def summary(records):
    files = [r for r in records if r.get("category") not in ("Folder", "Error")]
    by_cat = defaultdict(lambda: [0, 0])
    by_ext = defaultdict(lambda: [0, 0])
    for r in files:
        by_cat[r["category"]][0] += 1
        by_cat[r["category"]][1] += r.get("size", 0)
        ext = r.get("extension") or "(none)"
        by_ext[ext][0] += 1
        by_ext[ext][1] += r.get("size", 0)
    return {
        "total_files": len(files),
        "total_size": sum(r.get("size", 0) for r in files),
        "hidden_files": sum(1 for r in files if r.get("hidden")),
        "errors": sum(1 for r in records if r.get("error")),
        "folders": len({r.get("folder") for r in files}),
        "by_category": sorted(by_cat.items(), key=lambda kv: -kv[1][1]),
        "by_extension": sorted(by_ext.items(), key=lambda kv: -kv[1][0]),
        "largest": sorted(files, key=lambda r: -r.get("size", 0))[:25],
    }


def export(path, records, columns, split_metadata=False, dup_groups=None, scan_root=""):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        _export_xlsx(path, records, columns, split_metadata, dup_groups, scan_root)
    elif ext == ".json":
        keep = set(columns)
        out = [{k: v for k, v in r.items() if k in keep} for r in records]
        with open(path, "w", encoding="utf-8") as f:
            json.dump({"scanned_folder": scan_root, "exported": datetime.now().isoformat(timespec="seconds"),
                       "summary": {k: v for k, v in summary(records).items() if k != "largest"},
                       "files": out,
                       "duplicates": [{"group": g["id"], "size": g["size"], "wasted": g["wasted"],
                                       "files": [x["path"] for x in g["files"]]} for g in (dup_groups or [])]},
                      f, indent=2, default=str)
    elif ext in (".html", ".htm"):
        _export_html(path, records, columns, dup_groups, scan_root)
    else:
        headers, rows = _rows(records, columns, split_metadata)
        with open(path, "w", newline="", encoding="utf-8-sig") as f:  # BOM so Excel reads accents correctly
            writer = csv.writer(f, delimiter="\t" if ext == ".tsv" else ",")
            writer.writerow(headers)
            writer.writerows(rows)


def export_duplicates_csv(path, dup_groups, keepers=None):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Group", "Suggested action", "File name", "Full path", "Size", "Size (bytes)", "Created",
                    "Modified", "Resolution", "Length", "Checksum", "Metadata"])
        for g in dup_groups:
            for r in g["files"]:
                keep = keepers and keepers.get(g["id"]) is r
                w.writerow([g["id"], "Keep" if keep else "Duplicate", r["name"], r["path"], r.get("size_human", ""),
                            r.get("size", ""), r.get("created", ""), r.get("modified", ""), r.get("resolution", ""),
                            r.get("duration_human", ""), r.get("hash", ""), _metadata_text(r.get("metadata"))])


def _export_xlsx(path, records, columns, split_metadata, dup_groups, scan_root):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="DDE6F0")
    wb = Workbook()
    ws = wb.active
    ws.title = "Files"
    headers, rows = _rows(records, columns, split_metadata)
    ws.append(headers)
    for row in rows:
        ws.append([str(v) if isinstance(v, (dict, list)) else v for v in row])
    for cell in ws[1]:
        cell.font, cell.fill = bold, fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, len(str(h)) + 4))

    s = summary(records)
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Scanned folder", scan_root])
    ws2.append(["Exported", datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws2.append(["Total files", s["total_files"]])
    ws2.append(["Total size", human_size(s["total_size"])])
    ws2.append(["Folders containing files", s["folders"]])
    ws2.append(["Hidden files", s["hidden_files"]])
    ws2.append([])
    ws2.append(["Category", "Files", "Size", "Bytes"])
    ws2[ws2.max_row][0].font = bold
    for cat, (n, size) in s["by_category"]:
        ws2.append([cat, n, human_size(size), size])
    ws2.append([])
    ws2.append(["Extension", "Files", "Size", "Bytes"])
    for ext, (n, size) in s["by_extension"]:
        ws2.append([ext, n, human_size(size), size])
    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 60

    if dup_groups:
        ws3 = wb.create_sheet("Duplicates")
        ws3.append(["Group", "File name", "Full path", "Size", "Created", "Modified", "Resolution", "Length", "Checksum"])
        for cell in ws3[1]:
            cell.font, cell.fill = bold, fill
        for g in dup_groups:
            for r in g["files"]:
                ws3.append([g["id"], r["name"], r["path"], r.get("size_human", ""), r.get("created", ""),
                            r.get("modified", ""), r.get("resolution", ""), r.get("duration_human", ""), r.get("hash", "")])
        ws3.freeze_panes = "A2"
        ws3.column_dimensions["C"].width = 80
    wb.save(path)


def _export_html(path, records, columns, dup_groups, scan_root):
    s = summary(records)
    esc = html.escape
    headers, rows = _rows(records, columns, False)
    parts = [f"""<!doctype html><html><head><meta charset="utf-8"><title>File report - {esc(scan_root)}</title>
<style>
body{{font:14px -apple-system,Segoe UI,sans-serif;margin:24px;color:#222}}
h1{{font-size:20px}} h2{{font-size:16px;margin-top:32px}}
table{{border-collapse:collapse;font-size:12px}} th,td{{border:1px solid #ddd;padding:4px 6px;text-align:left;vertical-align:top}}
th{{background:#eef2f6;position:sticky;top:0}} tr:nth-child(even) td{{background:#fafafa}}
.cards{{display:flex;gap:12px;flex-wrap:wrap}} .card{{border:1px solid #ddd;border-radius:6px;padding:10px 14px}}
.card b{{display:block;font-size:20px}} input{{padding:6px;width:320px;margin:8px 0}}
@media (prefers-color-scheme:dark){{body{{background:#1b1d20;color:#ddd}}th{{background:#2a3038}}td{{border-color:#333}}tr:nth-child(even) td{{background:#202327}}}}
</style></head><body>
<h1>File report: {esc(scan_root)}</h1><p>Generated {datetime.now():%Y-%m-%d %H:%M}</p>
<div class="cards"><div class="card"><b>{s['total_files']:,}</b>files</div>
<div class="card"><b>{human_size(s['total_size'])}</b>total size</div>
<div class="card"><b>{s['hidden_files']:,}</b>hidden</div>
<div class="card"><b>{len(dup_groups or [])}</b>duplicate groups</div></div>
<h2>By category</h2><table><tr><th>Category</th><th>Files</th><th>Size</th></tr>"""]
    for cat, (n, size) in s["by_category"]:
        parts.append(f"<tr><td>{esc(cat)}</td><td>{n:,}</td><td>{human_size(size)}</td></tr>")
    parts.append("</table>")
    if dup_groups:
        parts.append("<h2>Duplicates</h2><table><tr><th>Group</th><th>Path</th><th>Size</th><th>Modified</th></tr>")
        for g in dup_groups:
            for r in g["files"]:
                parts.append(f"<tr><td>{g['id']}</td><td>{esc(r['path'])}</td><td>{esc(r.get('size_human',''))}</td>"
                             f"<td>{esc(r.get('modified',''))}</td></tr>")
        parts.append("</table>")
    parts.append('<h2>All files</h2><input id="q" placeholder="Filter..." oninput="f()"><table id="t"><tr>')
    parts.extend(f"<th>{esc(str(h))}</th>" for h in headers)
    parts.append("</tr>")
    for row in rows:
        parts.append("<tr>" + "".join(f"<td>{esc(str(v))}</td>" for v in row) + "</tr>")
    parts.append("""</table><script>function f(){var q=document.getElementById('q').value.toLowerCase();
document.querySelectorAll('#t tr').forEach(function(r,i){if(i)r.style.display=r.textContent.toLowerCase().includes(q)?'':'none'})}</script>
</body></html>""")
    with open(path, "w", encoding="utf-8") as f:
        f.write("".join(parts))
